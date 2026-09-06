"""4-dimension analysis workbook tests."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.datamind.pipeline import run_pipeline
from src.datamind.selected_analysis import (
    build_template_analysis,
    load_selected_tables,
    write_template_analysis_report,
)

SAMPLE_DIR = Path(__file__).resolve().parents[1] / "data" / "sample_xbrl"


def _build(tmp_path: Path):
    extracted = tmp_path / "extracted.xlsx"
    run_pipeline(input_dir=SAMPLE_DIR, output_excel=extracted, input_url=None)
    tables = load_selected_tables(extracted)
    analysis = build_template_analysis(tables)
    return analysis


def test_analysis_workbook_has_four_sheets(tmp_path: Path):
    analysis = _build(tmp_path)
    out = tmp_path / "analysis_4dims.xlsx"
    write_template_analysis_report(out, analysis)

    sheets = list(pd.ExcelFile(out).sheet_names)
    assert sheets == [
        "维度1_管理规模",
        "维度2_业绩表现",
        "维度3_资产配置",
        "维度4_竞品对标",
    ]


def test_dimension4_has_five_charts_and_no_financial_health_sheet(tmp_path: Path):
    analysis = _build(tmp_path)
    out = tmp_path / "analysis_4dims.xlsx"
    write_template_analysis_report(out, analysis)

    wb = load_workbook(out)
    assert len(wb["维度4_竞品对标"]._charts) == 5
    assert "维度5_财务健康" not in wb.sheetnames


def test_missing_sheets_raise_clear_error(tmp_path: Path):
    empty = tmp_path / "empty.xlsx"
    pd.DataFrame({"a": [1]}).to_excel(empty, sheet_name="unrelated", index=False)
    try:
        load_selected_tables(empty)
    except ValueError as ex:
        assert "missing required sheets" in str(ex)
    else:
        raise AssertionError("expected ValueError for workbook without required sheets")
