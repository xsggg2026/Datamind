from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Dict

import pandas as pd
from openpyxl.chart import BarChart, LineChart, PieChart, Reference


def _add_bar_chart(ws, title: str, data_col: int, category_col: int, min_row: int, max_row: int, anchor: str):
    chart = BarChart()
    chart.title = title
    data = Reference(ws, min_col=data_col, min_row=min_row - 1, max_row=max_row)
    categories = Reference(ws, min_col=category_col, min_row=min_row, max_row=max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 7
    chart.width = 12
    ws.add_chart(chart, anchor)


def _add_line_chart(ws, title: str, data_col: int, category_col: int, min_row: int, max_row: int, anchor: str):
    chart = LineChart()
    chart.title = title
    data = Reference(ws, min_col=data_col, min_row=min_row - 1, max_row=max_row)
    categories = Reference(ws, min_col=category_col, min_row=min_row, max_row=max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 7
    chart.width = 12
    ws.add_chart(chart, anchor)


def _add_pie_chart(ws, title: str, data_col: int, category_col: int, min_row: int, max_row: int, anchor: str):
    chart = PieChart()
    chart.title = title
    data = Reference(ws, min_col=data_col, min_row=min_row - 1, max_row=max_row)
    categories = Reference(ws, min_col=category_col, min_row=min_row, max_row=max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 7
    chart.width = 10
    ws.add_chart(chart, anchor)


def write_excel_report(
    output_path: Path,
    cleaned_df: pd.DataFrame,
    tables: Dict[str, pd.DataFrame],
    extra_tables: Dict[str, pd.DataFrame] | None = None,
    include_analysis_sheets: bool = True,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        if include_analysis_sheets:
            cleaned_df.to_excel(writer, sheet_name="cleaned_data", index=False)
            tables["scale_ranking"].to_excel(writer, sheet_name="scale_ranking", index=False)
            tables["mom_change"].to_excel(writer, sheet_name="mom_change", index=False)
            tables["type_distribution"].to_excel(writer, sheet_name="type_distribution", index=False)
            tables["strategy_performance"].to_excel(writer, sheet_name="strategy_performance", index=False)
            tables["quarterly_company_board"].to_excel(writer, sheet_name="quarterly_board", index=False)
            tables["product_type_trend"].to_excel(writer, sheet_name="type_trend", index=False)
            tables["company_nav_trend"].to_excel(writer, sheet_name="company_trend", index=False)
        if extra_tables:
            for sheet_name, table in extra_tables.items():
                safe_name = sheet_name[:31]
                table.to_excel(writer, sheet_name=safe_name, index=False)
        if include_analysis_sheets:
            summary = pd.DataFrame(
                {
                    "metric": [
                        "total_funds",
                        "total_companies",
                        "total_nav",
                        "avg_growth_rate",
                        "generated_at",
                    ],
                    "value": [
                        int(cleaned_df["fund_name"].nunique()) if not cleaned_df.empty else 0,
                        int(cleaned_df["company"].nunique()) if not cleaned_df.empty else 0,
                        float(cleaned_df["fund_nav_total"].sum()) if not cleaned_df.empty else 0.0,
                        float(cleaned_df["unit_nav_growth_rate"].mean()) if not cleaned_df.empty else 0.0,
                        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    ],
                }
            )
            summary.to_excel(writer, sheet_name="summary", index=False)

            wb = writer.book

            ws_scale = wb["scale_ranking"]
            max_row_scale = max(ws_scale.max_row, 2)
            _add_bar_chart(
                ws_scale,
                title="Fund Scale Ranking",
                data_col=5,
                category_col=2,
                min_row=2,
                max_row=max_row_scale,
                anchor="H2",
            )

            ws_mom = wb["mom_change"]
            max_row_mom = max(ws_mom.max_row, 2)
            _add_line_chart(
                ws_mom,
                title="NAV MoM Change",
                data_col=6,
                category_col=3,
                min_row=2,
                max_row=max_row_mom,
                anchor="H2",
            )

            ws_type = wb["type_distribution"]
            max_row_type = max(ws_type.max_row, 2)
            _add_pie_chart(
                ws_type,
                title="Type Distribution by NAV",
                data_col=3,
                category_col=1,
                min_row=2,
                max_row=max_row_type,
                anchor="E2",
            )

            ws_qboard = wb["quarterly_board"]
            max_row_qboard = max(ws_qboard.max_row, 2)
            _add_line_chart(
                ws_qboard,
                title="Quarterly NAV QoQ Change",
                data_col=6,
                category_col=2,
                min_row=2,
                max_row=max_row_qboard,
                anchor="H2",
            )

            ws_company = wb["company_trend"]
            if ws_company.max_column >= 2 and ws_company.max_row >= 2:
                line = LineChart()
                line.title = "Company NAV Trend"
                data = Reference(ws_company, min_col=2, max_col=ws_company.max_column, min_row=1, max_row=ws_company.max_row)
                categories = Reference(ws_company, min_col=1, min_row=2, max_row=ws_company.max_row)
                line.add_data(data, titles_from_data=True)
                line.set_categories(categories)
                line.height = 7
                line.width = 13
                ws_company.add_chart(line, "H2")


def write_markdown_brief(output_path: Path, tables: Dict[str, pd.DataFrame]) -> None:
    def dataframe_to_markdown(df: pd.DataFrame) -> str:
        if df.empty:
            return "No data"
        headers = [str(col) for col in df.columns]
        lines = [
            "| " + " | ".join(headers) + " |",
            "| " + " | ".join(["---"] * len(headers)) + " |",
        ]
        for row in df.itertuples(index=False):
            values = [str(value) for value in row]
            lines.append("| " + " | ".join(values) + " |")
        return "\n".join(lines)

    scale_table = tables["scale_ranking"].head(10)
    perf_table = tables["strategy_performance"].head(10)

    lines = [
        "# Public Fund Competitor Report",
        "",
        f"Generated at: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        "",
        "## Top 10 Scale Ranking",
        "",
        dataframe_to_markdown(scale_table),
        "",
        "## Strategy Performance Snapshot",
        "",
        dataframe_to_markdown(perf_table),
        "",
    ]

    output_path.write_text("\n".join(lines), encoding="utf-8")
