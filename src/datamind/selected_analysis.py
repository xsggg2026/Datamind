from __future__ import annotations

from pathlib import Path
import re
from typing import Dict, List

import pandas as pd
from openpyxl.chart import BarChart, LineChart, PieChart, ScatterChart, RadarChart, BubbleChart, Reference, Series


MISSING_TOKEN = "missing"


FIELD_NAME_BY_NO = {
    1: "基金名称",
    7: "基金管理人",
    10: "本期已实现收益",
    11: "本期利润",
    13: "本期基金份额净值增长率",
    14: "期末基金资产净值",
    18: "过去3个月净值增长率",
    19: "过去6个月净值增长率",
    20: "过去1年净值增长率",
    21: "过去3年净值增长率",
    22: "过去5年净值增长率",
    25: "同期业绩比较基准收益率",
    52: "管理人报酬",
    53: "托管费",
    54: "交易费用",
    55: "其他费用",
    57: "期初基金净值",
    58: "本期申购",
    59: "本期赎回",
    60: "本期利润分配",
    61: "期末基金净值",
    62: "权益投资（股票）金额",
    63: "权益投资占净值比例",
    64: "固定收益投资（债券）金额",
    65: "固定收益投资占净值比例",
    66: "基金投资金额",
    67: "基金投资占净值比例",
    68: "买入返售金融资产金额",
    69: "买入返售占净值比例",
    70: "银行存款和结算备付金合计",
    71: "银行存款和结算备付金占比",
    72: "其他资产金额",
    73: "其他资产占比",
}


def _col_name(df: pd.DataFrame, candidates: List[str]) -> str:
    for c in candidates:
        if c in df.columns:
            return c
    raise KeyError(f"Missing expected column, candidates={candidates}")


def _label_col(df: pd.DataFrame) -> str:
    for c in df.columns:
        cs = str(c)
        if cs.endswith("项") or cs.endswith("项目") or cs.endswith("名称") or cs == "字段名":
            return cs
    raise KeyError("Cannot detect metric label column")


def _value_col(df: pd.DataFrame) -> str:
    if "数值" in df.columns:
        return "数值"
    for c in df.columns:
        if str(c).endswith("数值") or str(c) == "value":
            return str(c)
    raise KeyError("Cannot detect metric value column")


def _to_float(value: object) -> float:
    text = str(value or "").strip()
    if text == "" or text.lower() in {"nan", MISSING_TOKEN, "missingb"}:
        return 0.0
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    is_percent = "%" in text
    text = text.replace("%", "").replace(",", "")
    text = re.sub(r"[^0-9.\-]", "", text)
    if text in {"", "-", "."}:
        return 0.0
    val = float(text)
    if negative:
        val = -val
    if is_percent:
        return val / 100.0
    return val


def _to_text(value: object) -> str:
    text = str(value or "").strip()
    if text == "" or text.lower() == "nan":
        return MISSING_TOKEN
    if text.lower() == "missingb":
        return MISSING_TOKEN
    return text


def _metric_value(metric_df: pd.DataFrame, field_no: int) -> str:
    val_col = _value_col(metric_df)
    if "字段序号" in metric_df.columns or "field_no" in metric_df.columns:
        no_col = _col_name(metric_df, ["字段序号", "field_no"])
        row = metric_df[metric_df[no_col] == field_no]
        if row.empty:
            return MISSING_TOKEN
        return _to_text(row.iloc[0][val_col])

    target_name = FIELD_NAME_BY_NO.get(field_no, "")
    if not target_name:
        return MISSING_TOKEN
    label_col = _label_col(metric_df)
    row = metric_df[metric_df[label_col].astype(str) == target_name]
    if row.empty:
        return MISSING_TOKEN
    return _to_text(row.iloc[0][val_col])


def _metric_name(metric_df: pd.DataFrame, field_no: int) -> str:
    if "字段序号" in metric_df.columns or "field_no" in metric_df.columns:
        no_col = _col_name(metric_df, ["字段序号", "field_no"])
        name_col = _col_name(metric_df, ["字段名", "field_name"])
        row = metric_df[metric_df[no_col] == field_no]
        if row.empty:
            return FIELD_NAME_BY_NO.get(field_no, f"field_{field_no}")
        return str(row.iloc[0][name_col])
    return FIELD_NAME_BY_NO.get(field_no, f"field_{field_no}")


def _metric_missing(metric_df: pd.DataFrame, field_no: int) -> bool:
    return _metric_value(metric_df, field_no) == MISSING_TOKEN


def _metric_float_or_none(metric_df: pd.DataFrame, field_no: int):
    value = _metric_value(metric_df, field_no)
    if value == MISSING_TOKEN:
        return None
    return _to_float(value)


def _safe_ratio(numerator: float, denominator: float) -> float:
    if denominator == 0:
        return 0.0
    return numerator / denominator


def _safe_ratio_token(numerator, denominator):
    if numerator is None or denominator is None or denominator == 0:
        return MISSING_TOKEN
    return numerator / denominator


def load_selected_tables(path: Path) -> Dict[str, pd.DataFrame]:
    required = [
        "table01_basic_info",
        "table02_fin_kpi",
        "table03_nav_perf",
        "table04_balance",
        "table05_income",
        "table06_net_asset",
        "table07_portfolio",
        "table08_top10",
        "table09_industry",
        "table10_holders",
        "table11_fundraising",
    ]
    alias = {
        "table01_basic_info": "表01_基本信息",
        "table02_fin_kpi": "表02_主要财务指标",
        "table03_nav_perf": "表03_净值表现",
        "table04_balance": "表04_资产负债表",
        "table05_income": "表05_利润表",
        "table06_net_asset": "表06_净资产变动表",
        "table07_portfolio": "表07_投资组合报告",
        "table08_top10": "表08_前十大持仓",
        "table09_industry": "表09_行业配置",
        "table10_holders": "表10_基金持有人结构",
        "table11_fundraising": "表11_新发与募集信息",
    }
    with pd.ExcelFile(path) as xls:
        out: Dict[str, pd.DataFrame] = {}
        missing = []
        for key in required:
            if key in xls.sheet_names:
                out[key] = pd.read_excel(xls, sheet_name=key)
                continue
            zh = alias[key]
            if zh in xls.sheet_names:
                out[key] = pd.read_excel(xls, sheet_name=zh)
                continue
            missing.append(key)
    if missing:
        raise ValueError(f"Input file missing required sheets: {missing}")
    return out


def _build_dimension1(tables: Dict[str, pd.DataFrame]) -> Dict[str, pd.DataFrame]:
    basic = tables["table01_basic_info"]
    fin_kpi = tables["table02_fin_kpi"]
    net_asset = tables["table06_net_asset"]

    fund_name = _metric_value(basic, 1)
    company = _metric_value(basic, 7)

    nav_start = _metric_float_or_none(net_asset, 57)
    nav_end = _metric_float_or_none(fin_kpi, 14)
    sub = _metric_float_or_none(net_asset, 58)
    red = _metric_float_or_none(net_asset, 59)
    profit = _metric_float_or_none(fin_kpi, 11)
    dividend = _metric_float_or_none(net_asset, 60)

    trend = pd.DataFrame(
        {
            "阶段": ["期初", "期末"],
            "基金资产净值": [nav_start if nav_start is not None else MISSING_TOKEN, nav_end if nav_end is not None else MISSING_TOKEN],
        }
    )

    compare = pd.DataFrame(
        {
            "公司": [company],
            "基金名称": [fund_name],
            "期末资产净值": [nav_end if nav_end is not None else MISSING_TOKEN],
        }
    )

    ranking = compare.copy()
    ranking["_sort"] = ranking["期末资产净值"].apply(lambda x: _to_float(x) if x != MISSING_TOKEN else -1)
    ranking = ranking.sort_values("_sort", ascending=False).drop(columns=["_sort"]).reset_index(drop=True)
    ranking["规模排名"] = ranking.index + 1

    total_nav = ranking["期末资产净值"].apply(_to_float).sum()
    ranking_sorted = ranking.sort_values("期末资产净值", key=lambda s: s.apply(_to_float), ascending=False)
    cr3 = _safe_ratio(ranking_sorted.head(3)["期末资产净值"].sum(), total_nav)
    cr5 = _safe_ratio(ranking_sorted.head(5)["期末资产净值"].sum(), total_nav)
    cr10 = _safe_ratio(ranking_sorted.head(10)["期末资产净值"].sum(), total_nav)
    concentration = pd.DataFrame(
        {
            "指标": ["CR3", "CR5", "CR10"],
            "数值": [cr3, cr5, cr10],
        }
    )

    can_explain = all(x is not None for x in [nav_start, nav_end, sub, red, profit, dividend])
    if can_explain:
        explained_end = nav_start + sub - red + profit - dividend
        residual = nav_end - explained_end
    else:
        residual = MISSING_TOKEN
    attribution = pd.DataFrame(
        {
            "因子": ["期初净值", "申购", "赎回", "本期利润", "分红", "归因残差", "期末净值(实际)"],
            "数值": [
                nav_start if nav_start is not None else MISSING_TOKEN,
                sub if sub is not None else MISSING_TOKEN,
                (-red) if red is not None else MISSING_TOKEN,
                profit if profit is not None else MISSING_TOKEN,
                (-dividend) if dividend is not None else MISSING_TOKEN,
                residual,
                nav_end if nav_end is not None else MISSING_TOKEN,
            ],
        }
    )

    return {
        "d1_scale_trend": trend,
        "d1_company_compare": compare,
        "d1_scale_ranking": ranking,
        "d1_scale_concentration": concentration,
        "d1_scale_attribution": attribution,
    }


def _build_dimension2(tables: Dict[str, pd.DataFrame]) -> Dict[str, pd.DataFrame]:
    nav_perf = tables["table03_nav_perf"]
    fin_kpi = tables["table02_fin_kpi"]

    stage_map = {
        18: "过去3个月",
        19: "过去6个月",
        20: "过去1年",
        21: "过去3年",
        22: "过去5年",
        23: "过去7年",
        24: "过去10年",
    }
    rows: List[Dict[str, object]] = []
    for no, stage in stage_map.items():
        raw_val = _metric_value(nav_perf, no)
        val = raw_val if raw_val == MISSING_TOKEN else _to_float(raw_val)
        rows.append({"阶段": stage, "净值增长率": val})
    stage_returns = pd.DataFrame(rows)

    bench_raw = _metric_value(nav_perf, 25)
    bench = bench_raw if bench_raw == MISSING_TOKEN else _to_float(bench_raw)
    excess = stage_returns.copy()
    excess["业绩比较基准收益率"] = bench
    def _calc_excess(row):
        if row["净值增长率"] == MISSING_TOKEN or row["业绩比较基准收益率"] == MISSING_TOKEN:
            return MISSING_TOKEN
        return row["净值增长率"] - row["业绩比较基准收益率"]
    excess["超额收益"] = excess.apply(_calc_excess, axis=1)

    one_year = stage_returns[stage_returns["阶段"] == "过去1年"]["净值增长率"].iloc[0]
    rank = pd.DataFrame(
        {
            "基金": [_metric_value(tables["table01_basic_info"], 1)],
            "同期收益率": [one_year],
            "排名": [1],
        }
    )

    scatter = pd.DataFrame(
        {
            "基金": [_metric_value(tables["table01_basic_info"], 1)],
            "收益率": [one_year],
            "波动率": [0.0],
            "说明": ["当前11表未包含标准差字段，波动率待补充"],
        }
    )

    cumulative = pd.DataFrame(
        {
            "对象": ["基金累计增长率", "基准累计收益率"],
            "数值": [
                _to_float(_metric_value(fin_kpi, 13)),
                bench if bench != MISSING_TOKEN else MISSING_TOKEN,
            ],
        }
    )

    return {
        "d2_stage_returns": stage_returns,
        "d2_excess_returns": excess,
        "d2_risk_return": scatter,
        "d2_peer_ranking": rank,
        "d2_cumulative_perf": cumulative,
    }


def _build_dimension3(tables: Dict[str, pd.DataFrame]) -> Dict[str, pd.DataFrame]:
    portfolio = tables["table07_portfolio"]
    top10 = tables["table08_top10"].copy()
    industry = tables["table09_industry"].copy()

    alloc_rows = [
        ("股票", _to_float(_metric_value(portfolio, 62)), _to_float(_metric_value(portfolio, 63))),
        ("债券", _to_float(_metric_value(portfolio, 64)), _to_float(_metric_value(portfolio, 65))),
        ("基金", _to_float(_metric_value(portfolio, 66)), _to_float(_metric_value(portfolio, 67))),
        ("买入返售", _to_float(_metric_value(portfolio, 68)), _to_float(_metric_value(portfolio, 69))),
        ("现金类", _to_float(_metric_value(portfolio, 70)), _to_float(_metric_value(portfolio, 71))),
        ("其他", _to_float(_metric_value(portfolio, 72)), _to_float(_metric_value(portfolio, 73))),
    ]
    alloc = pd.DataFrame(alloc_rows, columns=["资产类别", "金额", "占比"])

    industry_out = pd.DataFrame(columns=["行业类别", "行业公允价值", "占净值比例"])
    if not industry.empty:
        industry_out = industry[["行业类别", "行业公允价值", "占净值比例"]].copy()
        industry_out["行业公允价值"] = industry_out["行业公允价值"].apply(_to_float)
        industry_out["占净值比例"] = industry_out["占净值比例"].apply(_to_float)

    top10_out = pd.DataFrame(columns=["持仓排名", "股票/债券代码", "股票/债券名称", "持仓数量", "公允价值", "占基金资产净值比例"])
    if not top10.empty:
        top10_out = top10.copy()
        if "占基金资产净值比例" in top10_out.columns:
            top10_out["占基金资产净值比例"] = top10_out["占基金资产净值比例"].apply(_to_float)

    concentration = pd.DataFrame(
        {
            "指标": ["前十大重仓占净值比例合计"],
            "数值": [top10_out["占基金资产净值比例"].sum() if "占基金资产净值比例" in top10_out.columns else 0.0],
        }
    )

    bond_type = pd.DataFrame(
        {
            "债券品种": [MISSING_TOKEN],
            "占净值比例": [0.0],
            "说明": ["当前11表未拆出债券品种明细，需补充债券品种表"],
        }
    )

    return {
        "d3_asset_structure": alloc,
        "d3_industry_distribution": industry_out,
        "d3_top10_holdings": top10_out,
        "d3_top10_concentration": concentration,
        "d3_bond_type": bond_type,
    }


def _build_dimension4(tables: Dict[str, pd.DataFrame]) -> Dict[str, pd.DataFrame]:
    income = tables["table05_income"]
    fin_kpi = tables["table02_fin_kpi"]
    net_asset = tables["table06_net_asset"]

    profit = _to_float(_metric_value(fin_kpi, 11))
    nav_start = _metric_float_or_none(net_asset, 57)
    nav_end = _metric_float_or_none(fin_kpi, 14)
    avg_nav = None
    if nav_start is not None and nav_end is not None:
        avg_nav = (nav_start + nav_end) / 2

    profitability = pd.DataFrame(
        {
            "指标": ["本期利润", "利润状态"],
            "数值": [profit if profit is not None else MISSING_TOKEN, "盈利" if (profit is not None and profit >= 0) else ("亏损" if profit is not None else MISSING_TOKEN)],
        }
    )

    fees = pd.DataFrame(
        {
            "费用项": [
                _metric_name(income, 52),
                _metric_name(income, 53),
                _metric_name(income, 54),
                _metric_name(income, 55),
            ],
            "金额": [
                _to_float(_metric_value(income, 52)),
                _to_float(_metric_value(income, 53)),
                _to_float(_metric_value(income, 54)),
                _to_float(_metric_value(income, 55)),
            ],
        }
    )
    fee_total = fees["金额"].sum()
    fees["占总费用比例"] = fees["金额"].apply(lambda x: _safe_ratio(x, fee_total))

    profit_to_nav = pd.DataFrame(
        {
            "指标": ["本期利润", "平均资产净值", "净收益/规模比"],
            "数值": [
                profit if profit is not None else MISSING_TOKEN,
                avg_nav if avg_nav is not None else MISSING_TOKEN,
                _safe_ratio_token(profit, avg_nav),
            ],
        }
    )

    sub = _metric_float_or_none(net_asset, 58)
    red = _metric_float_or_none(net_asset, 59)
    net_flow = pd.DataFrame(
        {
            "指标": ["本期申购", "本期赎回", "申赎净额"],
            "数值": [
                sub if sub is not None else MISSING_TOKEN,
                red if red is not None else MISSING_TOKEN,
                (sub - red) if (sub is not None and red is not None) else MISSING_TOKEN,
            ],
        }
    )

    return {
        "d4_profitability": profitability,
        "d4_fee_structure": fees,
        "d4_profit_to_scale": profit_to_nav,
        "d4_net_flow": net_flow,
    }


def _build_dimension5(tables: Dict[str, pd.DataFrame]) -> Dict[str, pd.DataFrame]:
    basic = tables["table01_basic_info"]
    fin_kpi = tables["table02_fin_kpi"]
    portfolio = tables["table07_portfolio"]

    company = _metric_value(basic, 7)
    fund = _metric_value(basic, 1)
    nav = _to_float(_metric_value(fin_kpi, 14))
    ret = _to_float(_metric_value(fin_kpi, 13))
    equity_ratio = _to_float(_metric_value(portfolio, 63))
    bond_ratio = _to_float(_metric_value(portfolio, 65))

    company_scale = pd.DataFrame(
        {
            "公司": [company],
            "公司总规模": [nav],
            "说明": ["当前样本仅1家公司，竞品对标需更多公司数据"],
        }
    )

    peer_perf = pd.DataFrame(
        {
            "基金": [fund],
            "收益率_过去1年": [_to_float(_metric_value(tables["table03_nav_perf"], 20))],
            "收益率_过去3年": [_to_float(_metric_value(tables["table03_nav_perf"], 21))],
            "收益率_过去5年": [_to_float(_metric_value(tables["table03_nav_perf"], 22))],
        }
    )

    style = pd.DataFrame(
        {
            "基金": [fund],
            "股票占比": [equity_ratio],
            "债券占比": [bond_ratio],
            "收益率": [ret],
            "规模": [nav],
        }
    )

    quarter_track = pd.DataFrame(
        {
            "对象": [fund],
            "期初规模": [_to_float(_metric_value(tables["table06_net_asset"], 57))],
            "期末规模": [nav],
            "变化率": [
                _safe_ratio(
                    nav - _to_float(_metric_value(tables["table06_net_asset"], 57)),
                    _to_float(_metric_value(tables["table06_net_asset"], 57)),
                )
            ],
        }
    )

    return {
        "d5_company_scale": company_scale,
        "d5_peer_performance": peer_perf,
        "d5_style_scatter": style,
        "d5_scale_perf_matrix": style.copy(),
        "d5_quarter_track": quarter_track,
    }


def build_template_analysis(tables: Dict[str, pd.DataFrame]) -> Dict[str, pd.DataFrame]:
    outputs: Dict[str, pd.DataFrame] = {}
    outputs.update(_build_dimension1(tables))
    outputs.update(_build_dimension2(tables))
    outputs.update(_build_dimension3(tables))
    outputs.update(_build_dimension4(tables))
    outputs.update(_build_dimension5(tables))
    return outputs


def _block_to_sheet(writer: pd.ExcelWriter, sheet_name: str, title: str, df: pd.DataFrame, start_row: int) -> Dict[str, int]:
    ws = writer.book[sheet_name]
    ws.cell(row=start_row, column=1, value=title)
    df.to_excel(writer, sheet_name=sheet_name, startrow=start_row, startcol=0, index=False)
    header_row = start_row + 1
    data_start = header_row + 1
    data_end = data_start + len(df) - 1
    return {
        "header_row": header_row,
        "data_start": data_start,
        "data_end": data_end,
        "col_count": len(df.columns),
    }


def _chart_line(ws, title: str, min_col: int, cats_col: int, min_row: int, max_row: int, anchor: str):
    if max_row < min_row:
        return
    c = LineChart()
    c.title = title
    data = Reference(ws, min_col=min_col, min_row=min_row - 1, max_row=max_row)
    cats = Reference(ws, min_col=cats_col, min_row=min_row, max_row=max_row)
    c.add_data(data, titles_from_data=True)
    c.set_categories(cats)
    c.height = 7
    c.width = 12
    ws.add_chart(c, anchor)


def _chart_bar(ws, title: str, min_col: int, max_col: int, cats_col: int, min_row: int, max_row: int, anchor: str, horizontal: bool = False, stacked: bool = False):
    if max_row < min_row:
        return
    c = BarChart()
    c.title = title
    if horizontal:
        c.type = "bar"
    if stacked:
        c.grouping = "stacked"
        c.overlap = 100
    data = Reference(ws, min_col=min_col, max_col=max_col, min_row=min_row - 1, max_row=max_row)
    cats = Reference(ws, min_col=cats_col, min_row=min_row, max_row=max_row)
    c.add_data(data, titles_from_data=True)
    c.set_categories(cats)
    c.height = 7
    c.width = 12
    ws.add_chart(c, anchor)


def _chart_pie(ws, title: str, data_col: int, cats_col: int, min_row: int, max_row: int, anchor: str):
    if max_row < min_row:
        return
    c = PieChart()
    c.title = title
    data = Reference(ws, min_col=data_col, min_row=min_row - 1, max_row=max_row)
    cats = Reference(ws, min_col=cats_col, min_row=min_row, max_row=max_row)
    c.add_data(data, titles_from_data=True)
    c.set_categories(cats)
    c.height = 7
    c.width = 10
    ws.add_chart(c, anchor)


def _chart_scatter(ws, title: str, x_col: int, y_col: int, min_row: int, max_row: int, anchor: str):
    if max_row < min_row:
        return
    c = ScatterChart()
    c.title = title
    xvalues = Reference(ws, min_col=x_col, min_row=min_row, max_row=max_row)
    yvalues = Reference(ws, min_col=y_col, min_row=min_row, max_row=max_row)
    c.series.append(Series(yvalues, xvalues, title_from_data=False))
    c.height = 7
    c.width = 12
    ws.add_chart(c, anchor)


def _chart_radar(ws, title: str, min_col: int, max_col: int, cats_col: int, min_row: int, max_row: int, anchor: str):
    if max_row < min_row:
        return
    c = RadarChart()
    c.title = title
    data = Reference(ws, min_col=min_col, max_col=max_col, min_row=min_row - 1, max_row=max_row)
    cats = Reference(ws, min_col=cats_col, min_row=min_row, max_row=max_row)
    c.add_data(data, titles_from_data=True)
    c.set_categories(cats)
    c.height = 7
    c.width = 12
    ws.add_chart(c, anchor)


def _chart_bubble(ws, title: str, x_col: int, y_col: int, size_col: int, min_row: int, max_row: int, anchor: str):
    if max_row < min_row:
        return
    c = BubbleChart()
    c.title = title
    xvalues = Reference(ws, min_col=x_col, min_row=min_row, max_row=max_row)
    yvalues = Reference(ws, min_col=y_col, min_row=min_row, max_row=max_row)
    zvalues = Reference(ws, min_col=size_col, min_row=min_row, max_row=max_row)
    c.series.append(Series(yvalues, xvalues, zvalues, title_from_data=False))
    c.height = 7
    c.width = 12
    ws.add_chart(c, anchor)


def write_template_analysis_report(output_path: Path, analysis_tables: Dict[str, pd.DataFrame]) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        wb = writer.book
        for s in ["维度1_管理规模", "维度2_业绩表现", "维度3_资产配置", "维度4_竞品对标"]:
            wb.create_sheet(s)

        # 维度1
        s1 = "维度1_管理规模"
        b1 = _block_to_sheet(writer, s1, "单只基金规模变化趋势", analysis_tables["d1_scale_trend"], 1)
        b2 = _block_to_sheet(writer, s1, "同一公司多只基金规模对比", analysis_tables["d1_company_compare"], 16)
        type_compare = analysis_tables["d1_company_compare"].copy()
        type_compare["基金类型"] = "同类产品"
        b3 = _block_to_sheet(writer, s1, "不同公司同类基金规模对比", type_compare[["公司", "基金名称", "期末资产净值"]], 31)
        b4 = _block_to_sheet(writer, s1, "规模排名", analysis_tables["d1_scale_ranking"], 46)
        b5 = _block_to_sheet(writer, s1, "规模集中度", analysis_tables["d1_scale_concentration"], 61)
        b6 = _block_to_sheet(writer, s1, "规模变动归因", analysis_tables["d1_scale_attribution"], 76)
        ws1 = wb[s1]
        _chart_line(ws1, "单只基金规模变化趋势", 2, 1, b1["data_start"], b1["data_end"], "E2")
        _chart_bar(ws1, "同一公司多只基金规模对比", 3, 3, 2, b2["data_start"], b2["data_end"], "E17")
        _chart_bar(ws1, "不同公司同类基金规模对比", 3, 3, 1, b3["data_start"], b3["data_end"], "E32")
        _chart_bar(ws1, "规模排名", 3, 3, 2, b4["data_start"], b4["data_end"], "E47", horizontal=True)
        _chart_pie(ws1, "规模集中度", 2, 1, b5["data_start"], b5["data_end"], "E62")
        _chart_bar(ws1, "规模变动归因", 2, 2, 1, b6["data_start"], b6["data_end"], "E77", stacked=True)

        # 维度2
        s2 = "维度2_业绩表现"
        b21 = _block_to_sheet(writer, s2, "各阶段收益率对比", analysis_tables["d2_stage_returns"], 1)
        b22 = _block_to_sheet(writer, s2, "超额收益分析", analysis_tables["d2_excess_returns"], 20)
        b23 = _block_to_sheet(writer, s2, "收益-风险综合评估", analysis_tables["d2_risk_return"], 39)
        b24 = _block_to_sheet(writer, s2, "同类基金业绩排名", analysis_tables["d2_peer_ranking"], 54)
        b25 = _block_to_sheet(writer, s2, "累计业绩走势", analysis_tables["d2_cumulative_perf"], 69)
        ws2 = wb[s2]
        _chart_bar(ws2, "各阶段收益率对比", 2, 2, 1, b21["data_start"], b21["data_end"], "E2")
        _chart_bar(ws2, "超额收益分析", 4, 4, 1, b22["data_start"], b22["data_end"], "G21")
        _chart_scatter(ws2, "收益-风险综合评估", 3, 2, b23["data_start"], b23["data_end"], "E40")
        _chart_bar(ws2, "同类基金业绩排名", 2, 2, 1, b24["data_start"], b24["data_end"], "E55", horizontal=True)

        # 维度3
        s3 = "维度3_资产配置"
        b31 = _block_to_sheet(writer, s3, "大类资产配置结构", analysis_tables["d3_asset_structure"], 1)
        b32 = _block_to_sheet(writer, s3, "行业配置分布", analysis_tables["d3_industry_distribution"], 20)
        b33 = _block_to_sheet(writer, s3, "前十大重仓股明细", analysis_tables["d3_top10_holdings"], 45)
        b34 = _block_to_sheet(writer, s3, "前十大重仓股集中度", analysis_tables["d3_top10_concentration"], 70)
        b35 = _block_to_sheet(writer, s3, "债券品种配置", analysis_tables["d3_bond_type"], 85)
        ws3 = wb[s3]
        _chart_pie(ws3, "大类资产配置结构", 3, 1, b31["data_start"], b31["data_end"], "F2")
        _chart_bar(ws3, "资产配置变化趋势(当前样本快照)", 3, 3, 1, b31["data_start"], b31["data_end"], "P2", stacked=True)
        _chart_bar(ws3, "行业配置分布", 3, 3, 1, b32["data_start"], b32["data_end"], "H21", horizontal=True)
        _chart_bar(ws3, "前十大重仓股占比", 6, 6, 1, b33["data_start"], b33["data_end"], "H46")

        # 维度4：竞品对标
        s4 = "维度4_竞品对标"
        b51 = _block_to_sheet(writer, s4, "头部公司规模对比", analysis_tables["d5_company_scale"], 1)
        b52 = _block_to_sheet(writer, s4, "同类产品业绩对比", analysis_tables["d5_peer_performance"], 20)
        b53 = _block_to_sheet(writer, s4, "资产配置风格对比", analysis_tables["d5_style_scatter"], 38)
        b54 = _block_to_sheet(writer, s4, "规模-业绩矩阵", analysis_tables["d5_scale_perf_matrix"], 56)
        b55 = _block_to_sheet(writer, s4, "季度变动追踪", analysis_tables["d5_quarter_track"], 74)
        ws5 = wb[s4]
        _chart_bar(ws5, "头部公司规模对比", 2, 2, 1, b51["data_start"], b51["data_end"], "F2")
        _chart_radar(ws5, "同类产品业绩对比", 2, 4, 1, b52["data_start"], b52["data_end"], "F21")
        _chart_scatter(ws5, "资产配置风格对比", 2, 4, b53["data_start"], b53["data_end"], "H39")
        _chart_bubble(ws5, "规模-业绩矩阵", 5, 4, 5, b54["data_start"], b54["data_end"], "H57")
        _chart_line(ws5, "季度变动追踪", 3, 1, b55["data_start"], b55["data_end"], "H75")
